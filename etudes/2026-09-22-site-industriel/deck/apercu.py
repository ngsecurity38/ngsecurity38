"""
Aperçu des diapositives, rendu à la main.

LibreOffice ne fonctionne pas ici : il faut bien regarder ce qu'on livre.
Ce rendu lit le .pptx forme par forme et le redessine — approximatif sur
le crénage, fidèle sur la composition, ce qui est exactement ce qu'un
contrôle visuel cherche : les chevauchements, les débordements, les vides.
"""
import sys, io
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

E = 100  # pixels par pouce
po = lambda v: (Emu(v).inches if v is not None else 0)
POL = '/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf'
police = lambda t, g: ImageFont.truetype(POL % ('-Bold' if g else ''), max(7, int(t * E / 72)))

def couleur(f, defaut=None):
    try:
        if f and f.type is not None and f.fore_color and f.fore_color.type is not None:
            return '#' + str(f.fore_color.rgb)
    except Exception:
        pass
    return defaut

pres = Presentation(sys.argv[1])
W, H = int(13.333 * E), int(7.5 * E)
for n, d in enumerate(pres.slides, 1):
    fond = '#FFFFFF'
    try:
        if d.background.fill.type is not None:
            fond = couleur(d.background.fill, '#FFFFFF')
    except Exception:
        pass
    im = Image.new('RGB', (W, H), fond); dr = ImageDraw.Draw(im)
    for f in d.shapes:
        x, y = po(f.left) * E, po(f.top) * E
        w, h = po(f.width) * E, po(f.height) * E
        if f.shape_type == 13:  # image
            try:
                ph = Image.open(io.BytesIO(f.image.blob)).convert('RGB')
                im.paste(ph.resize((max(1,int(w)), max(1,int(h)))), (int(x), int(y)))
            except Exception:
                dr.rectangle([x, y, x+w, y+h], outline='#999')
            continue
        if not getattr(f, 'has_text_frame', False) or not f.text_frame.text.strip():
            # Les tableaux sont des GraphicFrame : on les esquisse en gris.
            if f.shape_type == 19 or not hasattr(f, 'fill'):
                dr.rectangle([x, y, x+w, y+h], outline='#C7CDD5')
                try:
                    for li, ligne in enumerate(f.table.rows):
                        yy = y + h * li / len(f.table.rows)
                        dr.line([x, yy, x+w, yy], fill='#E3E7EC')
                        cx = x
                        for ci, cell in enumerate(f.table.columns):
                            dr.text((cx+4, yy+3), f.table.cell(li, ci).text[:22],
                                    font=police(9, li == 0), fill='#1A1D23')
                            cx += po(cell.width) * E
                except Exception:
                    pass
                continue
            c = couleur(f.fill, None)
            if c:
                if 'ellipse' in str(f.shape_type).lower() or (f.shape_type == 1 and abs(w-h) < 3):
                    dr.ellipse([x, y, x+w, y+h], fill=c)
                else:
                    dr.rounded_rectangle([x, y, x+w, y+h], radius=6, fill=c, outline='#DDE1E7')
            continue
        c = couleur(f.fill, None)
        if c:
            dr.ellipse([x, y, x+w, y+h], fill=c) if abs(w-h) < 4 else \
                dr.rounded_rectangle([x, y, x+w, y+h], radius=6, fill=c)
        cy = y
        for p in f.text_frame.paragraphs:
            runs = p.runs or []
            txt = ''.join(r.text for r in runs) or p.text
            if not txt:
                cy += 10; continue
            t = max([r.font.size.pt for r in runs if r.font.size] or [18])
            g = any(r.font.bold for r in runs)
            col = None
            for r in runs:
                try:
                    if r.font.color and r.font.color.rgb: col = '#' + str(r.font.color.rgb); break
                except Exception: pass
            pol = police(t, g)
            mots, ligne = txt.split(' '), ''
            for mot in mots:
                essai = (ligne + ' ' + mot).strip()
                if dr.textlength(essai, font=pol) > w and ligne:
                    dr.text((x, cy), ligne, font=pol, fill=col or '#1A1D23')
                    cy += t * 1.22 * E / 72; ligne = mot
                else:
                    ligne = essai
            if ligne:
                dr.text((x, cy), ligne, font=pol, fill=col or '#1A1D23')
                cy += t * 1.3 * E / 72
    im.save(f'apercu-{n:02d}.png')
print(f'{len(pres.slides._sldIdLst)} aperçus écrits')
